from pathlib import Path

from docx import Document


docx_path = (
    Path.home()
    / "Desktop"
    / "\u4e2d\u671f"
    / "\u4e2d\u671f0421.2_\u9010\u9875\u53c2\u8003\u6587\u732e\u6574\u7406_v2.docx"
)

doc = Document(docx_path)
print(doc.paragraphs[0].text)
print("tables", len(doc.tables), "rows", len(doc.tables[0].rows) if doc.tables else 0)
for idx in [1, 8, 12]:
    row = doc.tables[0].rows[idx]
    print("|".join(cell.text.replace("\n", "/")[:120] for cell in row.cells))
