from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, Inches


out = Path.home() / "Desktop" / "\u4e2d\u671f" / "\u4e2d\u671f0421.2_\u9010\u9875\u53c2\u8003\u6587\u732e\u6574\u7406_v2.docx"

refs = {
    "R1": "Habibi, M., Foroughi, S., Karamzadeh, V. et al. Direct sound printing. Nature Communications 13, 1800 (2022). https://doi.org/10.1038/s41467-022-29395-1",
    "R2": "Derayatifar, M., Habibi, M., Bhat, R. et al. Holographic direct sound printing. Nature Communications 15, 6691 (2024). https://doi.org/10.1038/s41467-024-50923-8",
    "R3": "Derayatifar, M., Habibi, M., Bhat, R. et al. Penalization and deep learning algorithms in Holographic Direct Sound Printing to improve print uniformity. Additive Manufacturing 105, 104782 (2025). https://doi.org/10.1016/j.addma.2025.104782",
    "R4": "Burstow, R., Andres, D., Jimenez, N. et al. Acoustic holography in biomedical applications. Physics in Medicine & Biology 70(6), 06TR01 (2025). https://doi.org/10.1088/1361-6560/adb89a",
    "R5": "Melde, K., Mark, A. G., Qiu, T. et al. Holograms for acoustics. Nature 537, 518-522 (2016). https://doi.org/10.1038/nature19755",
    "R6": "Gerchberg, R. W., Saxton, W. O. A practical algorithm for the determination of phase from image and diffraction plane pictures. Optik 35, 237-246 (1972).",
    "R7": "Zeng, X., McGough, R. J. Optimal simulations of ultrasonic fields produced by large thermal therapy arrays using the angular spectrum approach. Journal of the Acoustical Society of America 125(5), 2967-2977 (2009). https://doi.org/10.1121/1.3097499",
    "R8": "Treeby, B. E., Cox, B. T. k-Wave: MATLAB toolbox for the simulation and reconstruction of photoacoustic wave fields. Journal of Biomedical Optics 15(2), 021314 (2010). https://doi.org/10.1117/1.3360308",
    "R9": "Melchels, F. P. W., Feijen, J., Grijpma, D. W. A review on stereolithography and its applications in biomedical engineering. Biomaterials 31(24), 6121-6130 (2010). https://doi.org/10.1016/j.biomaterials.2010.04.050",
    "R10": "Tumbleston, J. R., Shirvanyants, D., Ermoshkin, N. et al. Continuous liquid interface production of 3D objects. Science 347(6228), 1349-1352 (2015). https://doi.org/10.1126/science.aaa2397",
    "R11": "Kelly, B. E., Bhattacharya, I., Heidari, H. et al. Volumetric additive manufacturing via tomographic reconstruction. Science 363(6431), 1075-1079 (2019). https://doi.org/10.1126/science.aau7114",
    "R12": "Zhang, M., Jin, B., Hua, Y. et al. Reconfigurable dynamic acoustic holography with acoustically transparent and programmable metamaterial. Nature Communications 16, 9126 (2025). https://doi.org/10.1038/s41467-025-64154-y",
    "R13": "Sapareto, S. A., Dewey, W. C. Thermal dose determination in cancer therapy. International Journal of Radiation Oncology Biology Physics 10(6), 787-800 (1984). https://doi.org/10.1016/0360-3016(84)90379-1",
    "R15": "Cummer, S. A., Christensen, J., Alu, A. Controlling sound with acoustic metamaterials. Nature Reviews Materials 1, 16001 (2016). https://doi.org/10.1038/natrevmats.2016.1",
}

slides = [
    {
        "page": "第1页 封面",
        "topic": "题目页；展示课题方向",
        "use": "可不放文献。若导师要求每页都有，可放总参考：HDSP + DSP。",
        "refs": ["R1", "R2"],
        "footer": "Habibi et al., 2022; Derayatifar et al., 2024",
    },
    {
        "page": "第2页 研究背景与任务回顾",
        "topic": "HDSP、超声穿透、生物安全、声空化热点、逐层/面打印优势",
        "use": "DSP支撑声空化/声化学聚合；HDSP支撑一次性截面投射；声全息综述支撑生物医学应用；SLA综述用于对比光固化局限。",
        "refs": ["R1", "R2", "R4", "R9"],
        "footer": "Habibi et al., 2022; Derayatifar et al., 2024; Burstow et al., 2025; Melchels et al., 2010",
    },
    {
        "page": "第3页 研究任务与总体链路",
        "topic": "目标图案、相位反演、厚度映射、k-Wave声场、固化评价的完整链路",
        "use": "声全息板的基本思想引用Melde；HDSP链路引用Derayatifar；k-Wave仿真引用Treeby；ASM传播引用Zeng。",
        "refs": ["R2", "R5", "R7", "R8"],
        "footer": "Melde et al., 2016; Derayatifar et al., 2024; Zeng & McGough, 2009; Treeby & Cox, 2010",
    },
    {
        "page": "第4页 项目演化总览",
        "topic": "六阶段代码迭代：MATLAB原型、厚度调试、IASA、热固化、Python混合优化、出口诊断",
        "use": "这是本课题内部工作进展页，引用不宜过多。用核心背景文献和本课题代码记录即可。",
        "refs": ["R2", "R3", "R8"],
        "footer": "Derayatifar et al., 2024; Derayatifar et al., 2025; Treeby & Cox, 2010; 本课题代码迭代记录",
    },
    {
        "page": "第5页 阶段1 MATLAB初级链路建立",
        "topic": "目标图案生成、IASA初版、相位到厚度、k-Wave正向仿真",
        "use": "GS/IASA相位恢复基础引用Gerchberg-Saxton；声全息板引用Melde；k-Wave引用Treeby。",
        "refs": ["R5", "R6", "R8"],
        "footer": "Gerchberg & Saxton, 1972; Melde et al., 2016; Treeby & Cox, 2010",
    },
    {
        "page": "第6页 阶段2 相位板构建与厚度映射",
        "topic": "0-2pi映射、unwrap、高斯平滑、wrapped phase、全局偏置、误差扩散量化",
        "use": "相位板制造与厚度编码引用Melde；HDSP相位板应用引用Derayatifar；声学材料/超材料背景可引用Cummer或Zhang。",
        "refs": ["R2", "R5", "R12", "R15"],
        "footer": "Melde et al., 2016; Derayatifar et al., 2024; Cummer et al., 2016; Zhang et al., 2025",
    },
    {
        "page": "第7页 阶段3 IASA与相位图计算",
        "topic": "加权IASA、zero-padding、乘性权重、PAPO-IASA",
        "use": "GS/IASA基础引用Gerchberg-Saxton；ASM引用Zeng；HDSP均匀性优化和penalization/DL对比引用Derayatifar 2025。",
        "refs": ["R3", "R6", "R7"],
        "footer": "Gerchberg & Saxton, 1972; Zeng & McGough, 2009; Derayatifar et al., 2025",
    },
    {
        "page": "第8页 阶段4 热-固化评价闭环",
        "topic": "声压阈值截断、温度阈值、Arrhenius热剂量、后续声空化",
        "use": "DSP/HDSP支撑声致固化和空化；热剂量判据引用Sapareto-Dewey；k-Wave声场热源来源引用Treeby；声空化未来方向引用Burstow。",
        "refs": ["R1", "R2", "R4", "R8", "R13"],
        "footer": "Habibi et al., 2022; Derayatifar et al., 2024; Treeby & Cox, 2010; Sapareto & Dewey, 1984; Burstow et al., 2025",
    },
    {
        "page": "第9页 阶段5 MATLAB+Python混合优化主线形成",
        "topic": "PAPO、ASM物理传播、损失函数、梯度下降、PAPO-IASA",
        "use": "PAPO与DL/penalization对比引用Derayatifar 2025；动态声全息中的PANN/DL思想引用Zhang 2025；ASM引用Zeng。",
        "refs": ["R3", "R7", "R12"],
        "footer": "Derayatifar et al., 2025; Zeng & McGough, 2009; Zhang et al., 2025",
    },
    {
        "page": "第10页 当前版本总览",
        "topic": "目标、PAPO-IASA、离散厚度、k-Wave、出口分析、热/剂量/固化指标",
        "use": "该页是综合链路图，需要放方法核心引用：HDSP、均匀性优化、k-Wave、热剂量。",
        "refs": ["R2", "R3", "R8", "R13"],
        "footer": "Derayatifar et al., 2024; Derayatifar et al., 2025; Treeby & Cox, 2010; Sapareto & Dewey, 1984",
    },
    {
        "page": "第11页 阶段6 后续增强与支线探索",
        "topic": "理论相位与实际调制差距、出口面幅度/相位诊断、厚度板调制规律",
        "use": "声全息板会通过材料厚度引入相位调制，引用Melde；声学超材料/动态可编程调制引用Cummer和Zhang；HDSP应用引用Derayatifar。",
        "refs": ["R2", "R5", "R12", "R15"],
        "footer": "Melde et al., 2016; Derayatifar et al., 2024; Cummer et al., 2016; Zhang et al., 2025",
    },
    {
        "page": "第12页 中期阶段总结",
        "topic": "已完成、问题、下一步：主线稳固、厚度调制规律、声空化固化预测、论文数据",
        "use": "总结页建议放总括性核心文献，覆盖DSP/HDSP、算法优化、声空化未来方向。",
        "refs": ["R1", "R2", "R3", "R4"],
        "footer": "Habibi et al., 2022; Derayatifar et al., 2024; Derayatifar et al., 2025; Burstow et al., 2025",
    },
]


def set_font(run, size=9, bold=False):
    run.font.name = "SimSun"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.55)
section.right_margin = Inches(0.55)

style = doc.styles["Normal"]
style.font.name = "SimSun"
style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
style.font.size = Pt(9)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("《中期0421.2.pptx》逐页参考文献整理")
set_font(r, 16, True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("原则：每页引用2-4条最直接支撑本页论点的文献；项目迭代页可同时注明“本课题代码迭代记录”。")
set_font(r, 9)

table = doc.add_table(rows=1, cols=5)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
headers = ["页码", "页面主题", "建议引用", "支撑理由", "页脚短引格式"]
for cell, header in zip(table.rows[0].cells, headers):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(header)
    set_font(r, 8, True)

for item in slides:
    row = table.add_row().cells
    values = [
        item["page"],
        item["topic"],
        "\n".join(item["refs"]),
        item["use"],
        item["footer"],
    ]
    for cell, value in zip(row, values):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        p = cell.paragraphs[0]
        p.text = ""
        r = p.add_run(value)
        set_font(r, 7.5)

doc.add_paragraph()
p = doc.add_paragraph()
r = p.add_run("完整参考文献表")
set_font(r, 13, True)

for key, text in refs.items():
    p = doc.add_paragraph()
    r = p.add_run(f"[{key}] {text}")
    set_font(r, 9)

doc.add_paragraph()
p = doc.add_paragraph()
r = p.add_run("备注")
set_font(r, 11, True)
for note in [
    "第1页如果不要求页脚，可不放文献；若必须每页都有，建议放DSP与HDSP两篇总参考。",
    "第4页、第10页、第12页属于本课题阶段总结/综合链路页，建议保留“本课题代码迭代记录”作为内部依据，不要只堆外文献。",
    "第8页必须明确声空化是后续计划，不应写成当前已实现模型；当前版本主要是热剂量/Arrhenius固化预测。",
]:
    p = doc.add_paragraph(style=None)
    r = p.add_run(note)
    set_font(r, 9)

out.parent.mkdir(parents=True, exist_ok=True)
doc.save(out)
print(out)
